"""Render an edition JSON to a ten-slide 16:9 PPTX briefing with speaker notes.

Colours and type pairing follow the edition's theme. Each slide carries the
page's takeaways, a visual panel (top-five list, a compact table, or a key
callout), clickable source chips, and detailed speaker notes with full URLs.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .schema import CITE_RE
from .themes import Theme

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.55)


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#").upper())


def _strip_cites(text: str) -> str:
    return CITE_RE.sub("", text).replace("  ", " ").strip()


def _box(slide, x, y, w, h, *, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, name=None):
    shp = slide.shapes.add_shape(shape, x, y, w, h)
    if name:
        shp.name = name
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _text(slide, x, y, w, h, runs, *, size=14, color="#000000", font="Calibri", bold=False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None, italic=False):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, opts) tuples."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        pieces = [(para, {})] if isinstance(para, str) else para
        for txt, opts in pieces:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", italic)
            f.name = opts.get("font", font)
            f.color.rgb = _rgb(opts.get("color", color))
            if opts.get("link"):
                r.hyperlink.address = opts["link"]
        gaps = [o["space_after"] for _, o in pieces if "space_after" in o]
        if gaps:
            p.space_after = Pt(gaps[0])
    return tb


def _bullets(slide, x, y, w, h, items, t: Theme, size=16):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.name = "Takeaways"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        n = p.add_run()
        n.text = f"{i + 1:02d}  "
        n.font.size = Pt(size)
        n.font.bold = True
        n.font.name = t.pptx_head
        n.font.color.rgb = _rgb(t.accent)
        r = p.add_run()
        r.text = _strip_cites(item)
        r.font.size = Pt(size)
        r.font.name = t.pptx_body
        r.font.color.rgb = _rgb(t.ink)
    return tb


def _panel_items(page: dict) -> tuple[str, list[str]] | None:
    for b in page["blocks"]:
        if b["type"] == "headlines" and b["items"]:
            return "Top five", b["items"][:5]
    return None


def _first_table(page: dict) -> dict | None:
    return next((b for b in page["blocks"] if b["type"] == "table" and b["rows"]), None)


def _first_callout(page: dict) -> dict | None:
    return next((b for b in page["blocks"] if b["type"] == "callout"), None)


def _source_chips(slide, edition, cites, t: Theme, y, dark=False):
    srcs = {s["id"]: s for s in edition["sources"]}
    ids = [c for c in dict.fromkeys(cites) if c in srcs][:7]
    if not ids:
        return
    # PowerPoint draws hyperlinks in the theme's link colour, so on dark slides the
    # chips sit on a light strip to keep them legible.
    head, link = t.primary, t.accent
    if dark:
        strip = _box(slide, MARGIN - Inches(0.12), y - Inches(0.06), W - 2 * MARGIN + Inches(0.24), Inches(0.38),
                     fill="#FFFFFF", shape=MSO_SHAPE.ROUNDED_RECTANGLE, name="Sources strip")
        strip.adjustments[0] = 0.3
    runs = [("Sources  ", {"bold": True, "color": head, "size": 10, "font": t.pptx_body})]
    for i, c in enumerate(ids):
        s = srcs[c]
        label = f"[{c}] {s['publisher']}"
        runs.append((label, {"size": 10, "color": link, "link": s["url"], "font": t.pptx_body}))
        if i < len(ids) - 1:
            runs.append(("   ", {"size": 10}))
    _text(slide, MARGIN, y, W - 2 * MARGIN, Inches(0.3), [runs], name="Sources")


def _all_cites(page: dict) -> list[int]:
    nums = list(page["slide"]["cites"])
    for b in page["blocks"]:
        nums += b["cites"]
        for txt in (b["what"], b["why"], b["impact"], *b["body"], *b["items"]):
            for m in CITE_RE.finditer(txt):
                nums += [int(n) for n in m.group(1).split(",")]
    return nums


def _notes(page: dict, edition: dict) -> str:
    srcs = {s["id"]: s for s in edition["sources"]}
    lines = [page["slide"]["notes"].strip(), ""]
    used = [c for c in dict.fromkeys(_all_cites(page)) if c in srcs]
    if used:
        lines.append("Sources cited on this page:")
        for c in used:
            s = srcs[c]
            lines.append(f"[{c}] {s['publisher']}: {s['title']} ({s['published']}) {s['url']}")
    return "\n".join(lines)


def render_pptx(edition: dict, theme: Theme, out: Path) -> Path:
    t = theme
    meta = edition["meta"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    content_w = W - 2 * MARGIN

    for page in edition["pages"]:
        s = prs.slides.add_slide(blank)
        dark = page["number"] in (1, 10)
        bg = t.primary if dark else t.paper
        fg = "#FFFFFF" if dark else t.ink
        head_col = "#FFFFFF" if dark else t.primary
        _box(s, 0, 0, W, H, fill=bg, name="Background")

        # header line: brand · section · date
        _text(s, MARGIN, Inches(0.32), content_w, Inches(0.3),
              [[(meta["title"], {"bold": True, "font": t.pptx_head, "color": head_col, "size": 12}),
                (f"   ·   {page['number']:02d}  {page['section'].upper()}", {"bold": True, "color": t.accent if not dark else "#9FD6D0", "size": 11}),
                (f"   ·   {meta['edition_date']}", {"color": fg, "size": 11})]],
              font=t.pptx_body, name="Running header")

        title = page["slide"]["title"] or page["headline"]
        _text(s, MARGIN, Inches(0.78), content_w, Inches(1.05), _strip_cites(title),
              size=34 if len(title) < 60 else 30, bold=True, font=t.pptx_head, color=head_col,
              anchor=MSO_ANCHOR.TOP, name="Title")

        top = Inches(2.0)
        body_h = Inches(4.55)
        left_w = Inches(7.3)
        gap = Inches(0.4)
        right_x = MARGIN + left_w + gap
        right_w = W - MARGIN - right_x

        if page["number"] == 1:
            _text(s, MARGIN, Inches(1.72), content_w, Inches(0.3),
                  f"Coverage window: {meta['window_start']} to {meta['window_end']} ({meta['timezone']})",
                  size=12, color="#CFE3E0", font=t.pptx_body, italic=True, name="Window")

        _bullets(s, MARGIN, top + Inches(0.15), left_w, body_h, page["slide"]["takeaways"][:5],
                 t, size=16 if sum(len(x) for x in page["slide"]["takeaways"]) < 520 else 14)
        if dark:  # re-colour bullet text for dark slides
            for p in s.shapes[-1].text_frame.paragraphs:
                for r in p.runs[1:]:
                    r.font.color.rgb = _rgb("#FFFFFF")

        # right-hand visual panel
        panel_fill = "#FFFFFF" if dark else t.accent_soft
        panel = _box(s, right_x, top, right_w, body_h, fill=panel_fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name="Panel")
        panel.adjustments[0] = 0.04
        px, pw = right_x + Inches(0.3), right_w - Inches(0.6)
        items = _panel_items(page)
        table = _first_table(page)
        callout = _first_callout(page)
        if items:
            label, lines = items
            _text(s, px, top + Inches(0.25), pw, Inches(0.3), label.upper(), size=11, bold=True, color=t.accent, font=t.pptx_body)
            paras = []
            for i, ln in enumerate(lines):
                paras.append([(f"{i + 1}  ", {"bold": True, "color": t.accent, "size": 14, "font": t.pptx_head}),
                              (_strip_cites(ln), {"color": t.primary, "size": 12, "font": t.pptx_body, "space_after": 8})])
            _text(s, px, top + Inches(0.65), pw, body_h - Inches(0.9), paras, name="Panel list")
        elif table:
            cols = table["columns"][:3]
            rows = [r[:3] for r in table["rows"][:6]]
            _text(s, px, top + Inches(0.25), pw, Inches(0.3), (table["title"] or "At a glance").upper(),
                  size=11, bold=True, color=t.accent, font=t.pptx_body)
            gt = s.shapes.add_table(len(rows) + 1, len(cols), px, top + Inches(0.65), pw, Inches(0.3) * (len(rows) + 1))
            gt.name = "Panel table"
            tbl = gt.table
            for ci, c in enumerate(cols):
                cell = tbl.cell(0, ci)
                cell.text = c
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(t.primary)
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size, r.font.bold, r.font.name = Pt(10), True, t.pptx_body
                        r.font.color.rgb = _rgb("#FFFFFF")
            for ri, row in enumerate(rows, 1):
                for ci, val in enumerate(row):
                    cell = tbl.cell(ri, ci)
                    cell.text = _strip_cites(val)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb("#FFFFFF" if ri % 2 else t.stripe)
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size, r.font.name = Pt(10), t.pptx_body
                            r.font.bold = ci == 0
                            r.font.color.rgb = _rgb(t.primary if ci == 0 else t.ink)
        else:
            src = callout or next((b for b in page["blocks"] if b["type"] == "story"), None)
            label = (src["kicker"] or "Key point") if src else "Key point"
            heading = _strip_cites(src["title"]) if src else ""
            text = _strip_cites((src["why"] or (src["body"][0] if src["body"] else "")) if src else "")
            if len(text) > 330:
                text = text[:330].rsplit(" ", 1)[0] + "…"
            _text(s, px, top + Inches(0.25), pw, Inches(0.3), label.upper(), size=11, bold=True, color=t.accent, font=t.pptx_body)
            _text(s, px, top + Inches(0.65), pw, Inches(1.2), heading, size=20, bold=True, color=t.primary, font=t.pptx_head)
            if src and src["items"]:
                paras = [[("•  ", {"bold": True, "color": t.accent, "size": 13}),
                          (_strip_cites(it), {"color": t.ink, "size": 13, "font": t.pptx_body, "space_after": 6})]
                         for it in src["items"][:4]]
                _text(s, px, top + Inches(1.85), pw, body_h - Inches(2.05), paras, name="Panel items")
            else:
                _text(s, px, top + Inches(1.95), pw, body_h - Inches(2.2), text, size=13, color=t.ink, font=t.pptx_body)

        _source_chips(s, edition, _all_cites(page), t, Inches(6.78), dark=dark)
        _text(s, W - MARGIN - Inches(1.2), Inches(7.08), Inches(1.2), Inches(0.25), f"{page['number']} / 10",
              size=10, color=fg, font=t.pptx_body, align=PP_ALIGN.RIGHT, name="Slide number")
        s.notes_slide.notes_text_frame.text = _notes(page, edition)

    prs.core_properties.title = f"{meta['title']} — {meta['edition_date']}"
    prs.core_properties.subject = f"Coverage window {meta['window_start']} to {meta['window_end']}"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
