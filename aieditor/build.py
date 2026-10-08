"""Build an edition: validate JSON, pick a theme, render PDF + PPTX, run QA.

    python -m aieditor.build editions/2026-10-08/edition.json
    python -m aieditor.build editions/2026-10-08/edition.json --theme markets
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path

from . import schema
from .render_pdf import render_pdf
from .render_pptx import render_pptx
from .themes import THEMES, choose_theme, score_themes

ROOT = Path(__file__).resolve().parent.parent
EDITIONS = Path(os.environ.get("EDITIONS_DIR", ROOT / "editions"))


def previous_theme(edition_dir: Path) -> str | None:
    """Theme of the most recent earlier edition, so consecutive days differ."""
    earlier = sorted(p for p in EDITIONS.glob("*/build.json") if p.parent.name < edition_dir.name)
    if not earlier:
        return None
    try:
        return json.loads(earlier[-1].read_text())["theme"]
    except (KeyError, ValueError):
        return None


def qa_pdf(pdf: Path) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(str(pdf))
    links = 0
    for page in reader.pages:
        for annot in page.get("/Annots") or []:
            a = annot.get_object()
            if a.get("/Subtype") == "/Link" and "/A" in a and "/URI" in a["/A"]:
                links += 1
    return {"pages": len(reader.pages), "links": links}


def qa_pptx(pptx: Path) -> dict:
    from pptx import Presentation

    prs = Presentation(str(pptx))
    notes = sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
    links = sum(
        1
        for s in prs.slides
        for sh in s.shapes
        if sh.has_text_frame
        for p in sh.text_frame.paragraphs
        for r in p.runs
        if r.hyperlink.address
    )
    ratio = prs.slide_width / prs.slide_height
    return {"slides": len(prs.slides), "with_notes": notes, "links": links, "aspect": round(ratio, 3)}


def build(edition_path: Path, theme_key: str | None = None) -> dict:
    edition = schema.load(edition_path)
    out_dir = edition_path.parent
    theme = choose_theme(edition, previous=previous_theme(out_dir), override=theme_key)
    stem = out_dir.name
    pdf = out_dir / f"rajesh-ai-it-daily-{stem}.pdf"
    pptx = out_dir / f"rajesh-ai-it-daily-{stem}.pptx"

    fit = render_pdf(edition, theme, pdf, html_out=out_dir / "edition.html")
    render_pptx(edition, theme, pptx)

    result = {
        "edition": stem,
        "theme": theme.key,
        "theme_name": theme.name,
        "theme_scores": {k: round(v, 2) for k, v in score_themes(edition).items()},
        "words": schema.word_count(edition),
        "pdf": {"file": pdf.name, **qa_pdf(pdf), "fit": fit},
        "pptx": {"file": pptx.name, **qa_pptx(pptx)},
    }
    problems = []
    if result["pdf"]["pages"] != 10:
        problems.append(f"PDF has {result['pdf']['pages']} pages")
    if result["pptx"]["slides"] != 10:
        problems.append(f"PPTX has {result['pptx']['slides']} slides")
    if result["pptx"]["with_notes"] != 10:
        problems.append("some slides lack speaker notes")
    if abs(result["pptx"]["aspect"] - 16 / 9) > 0.01:
        problems.append("PPTX is not 16:9")
    thin = [f["page"] for f in fit if f["fill"] < 0.6]
    if thin:
        problems.append(f"pages look underfilled: {thin}")
    result["problems"] = problems
    (out_dir / "build.json").write_text(json.dumps(result, indent=2))
    return result


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("edition", type=Path, help="path to edition.json")
    ap.add_argument("--theme", choices=sorted(THEMES), help="force a theme instead of choosing from content")
    args = ap.parse_args(argv)
    result = build(args.edition.resolve(), args.theme)
    print(json.dumps(result, indent=2))
    return 1 if result["problems"] else 0


if __name__ == "__main__":
    sys.exit(main())
